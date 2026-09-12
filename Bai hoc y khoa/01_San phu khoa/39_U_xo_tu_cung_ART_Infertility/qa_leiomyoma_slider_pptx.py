# -*- coding: utf-8 -*-
"""QA gate for DOCX-only leiomyoma ART slider3636 deck."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.stderr.reconfigure(encoding="utf-8", errors="replace")

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).parent
DATE = "2026-07-13"
PPTX = ROOT / f"U_xo_tu_cung_leiomyoma_slider3636_long_smooth_{DATE}.pptx"
DECK = ROOT / f"U_xo_tu_cung_leiomyoma_slider3636_long_{DATE}.deck.json"
PROVENANCE = ROOT / f"U_xo_tu_cung_leiomyoma_slider3636_long_{DATE}.provenance.json"
APPROVED_DOCX = ROOT / "U_xo_tu_cung_ART_Infertility_2026-07-12.docx"
EXPECTED_SLIDES = 104
REQUIRED_TABLES = {"T3", "T5", "T11", "T15", "T20", "T21", "T24"}
FORBIDDEN_TEXT = [
    "Quá trình mổ u xơ", "Giải thích các bước phẫu thuật u xơ",
    "Hạn chế dính buồng tử cung", "PMC13278667_fulltext",
    "placeholder", "lorem", "ipsum", "xxxx",
    "Làm gì:", "Tại sao:", "Nếu bỏ qua/làm sai:", "Nguy cơ nếu sai:",
]
# Visible English phrases that must not appear in rendered PPTX text.
# Allowed acronyms/medical terms: ART, IVF, FET, FIGO, TVUS, SIS, MRI, RIF,
# AMH, AFC, AUB, Hb, GnRH, UAE, HIFU, RFA, NSAID, TXA, STUMP, DOR, PMID, DOI.
FORBIDDEN_ENGLISH_VISIBLE = [
    "Roadmap", "ART risk", "cavity-distorting", "restore cavity", "gray zone",
    "Shared-decision checklist", "Management by FIGO", "Treatment-effect matrix",
    "Embryo banking decision", "Bank embryos first", "Operate first",
    "Fresh vs frozen", "Timing algorithm", "Integrated ART decision algorithm",
    "Take-home messages", "Case drill", "Decision by FIGO",
    "Mechanism → decision", "prove the cavity", "usually not just for IVF",
    "shared decision", "anchoring bias", "over-treatment",
]
REQUIRED_SECTION_TITLES = [
    "Khung tư duy lâm sàng", "Mô tả tổn thương và FIGO", "Phân tầng nguy cơ ART",
    "Cơ chế ảnh hưởng sinh sản", "Chẩn đoán để trả lời câu hỏi ART",
    "Đích điều trị trước khi chọn phương pháp", "Xử trí theo FIGO và bối cảnh ART",
    "Phẫu thuật và cái giá", "Thuốc và can thiệp không phẫu thuật",
    "Tạo và trữ phôi", "Theo dõi và tư vấn thai kỳ", "Tích hợp và kết thúc",
]
MAX_VISIBLE_CHARS = 1200
MAX_DENSE_SLIDES = 12
DENSE_CHARS = 950
MIN_VISIBLE_CHARS = 90  # slides with too little visible content


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def shape_text(shape):
    parts = []
    if getattr(shape, "has_text_frame", False):
        parts.append(shape.text_frame.text or "")
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                parts.append(cell.text or "")
    return "\n".join(p for p in parts if p)


def slide_text(slide):
    return "\n".join(shape_text(s) for s in iter_shapes(slide.shapes) if shape_text(s))


def notes_text(slide):
    if not slide.has_notes_slide:
        return ""
    return (slide.notes_slide.notes_text_frame.text or "").strip()


def fail(msg: str, failures: list[str]):
    print(f"[FAIL] {msg}")
    failures.append(msg)


def ok(msg: str):
    print(f"[OK] {msg}")


def picture_count(prs: Presentation) -> int:
    total = 0
    for slide in prs.slides:
        for s in iter_shapes(slide.shapes):
            if s.shape_type == MSO_SHAPE_TYPE.PICTURE:
                total += 1
    return total


def page_badges(text: str):
    # slider3636 page badge is formatted with spaces around slash: "34  /  91".
    # Avoid false positives like "type 3/4" inside medical content.
    return re.findall(r"\b(\d+)\s+/\s+(\d+)\b", text)


def main():
    failures: list[str] = []
    if not PPTX.exists():
        raise SystemExit(f"Missing PPTX: {PPTX}")
    if not DECK.exists():
        raise SystemExit(f"Missing deck JSON: {DECK}")
    if not PROVENANCE.exists():
        raise SystemExit(f"Missing provenance JSON: {PROVENANCE}")

    deck = json.loads(DECK.read_text(encoding="utf-8"))
    prov = json.loads(PROVENANCE.read_text(encoding="utf-8"))
    slides = deck.get("slides", [])
    prov_slides = prov.get("slides", [])
    prs = Presentation(str(PPTX))
    slide_texts = [slide_text(slide).strip() for slide in prs.slides]
    all_text = "\n".join(slide_texts)
    lower = all_text.lower()

    print(f"PPTX: {PPTX}")
    print(f"slides: {len(prs.slides)}")
    print(f"deck_slides: {len(slides)}")
    print(f"provenance_entries: {len(prov_slides)}")

    if len(prs.slides) == EXPECTED_SLIDES and len(slides) == EXPECTED_SLIDES and len(prov_slides) == EXPECTED_SLIDES:
        ok(f"counts match expected {EXPECTED_SLIDES}")
    else:
        fail(f"count mismatch pptx={len(prs.slides)} deck={len(slides)} prov={len(prov_slides)} expected={EXPECTED_SLIDES}", failures)

    src = Path(prov.get("source_docx", ""))
    if src.resolve() == APPROVED_DOCX.resolve():
        ok("approved DOCX source path")
    else:
        fail(f"wrong provenance source {src}", failures)

    wrong_sources = []
    for e in prov_slides:
        if Path(e.get("source_docx", "")).resolve() != APPROVED_DOCX.resolve():
            wrong_sources.append(e.get("slide_index"))
    if not wrong_sources:
        ok("all manifest entries use approved DOCX")
    else:
        fail(f"slides with wrong source: {wrong_sources[:20]}", failures)

    used_ids = {sid for e in prov_slides for sid in e.get("source_ids", [])}
    if REQUIRED_TABLES <= used_ids:
        ok(f"required backbone tables covered: {sorted(REQUIRED_TABLES)}")
    else:
        fail(f"missing backbone tables: {sorted(REQUIRED_TABLES - used_ids)}", failures)

    forbidden_hits = [x for x in FORBIDDEN_TEXT if x.lower() in lower or x in all_text]
    if not forbidden_hits:
        ok("no forbidden/source-leak/visible-rationale text")
    else:
        fail(f"forbidden text hits: {forbidden_hits}", failures)

    english_hits = [x for x in FORBIDDEN_ENGLISH_VISIBLE if x in all_text]
    if not english_hits:
        ok("no visible English leakage")
    else:
        fail(f"visible English leakage: {english_hits}", failures)

    pics = picture_count(prs)
    if pics == 0:
        ok("no embedded factual images")
    else:
        fail(f"embedded picture count should be 0, got {pics}", failures)

    # Notes/provenance.
    thin_notes = []
    missing_source_notes = []
    for i, slide in enumerate(prs.slides, 1):
        n = notes_text(slide)
        stype = slides[i-1].get("type") if i <= len(slides) else ""
        if i > len(slides):
            continue
        if len(n) < 30:
            thin_notes.append(i)
        if stype not in {"title", "outline", "section"} and "Nguồn DOCX:" not in n:
            missing_source_notes.append(i)
    if not thin_notes and not missing_source_notes:
        ok("speaker notes present with DOCX source IDs")
    else:
        if thin_notes:
            fail(f"thin notes: {thin_notes[:20]}", failures)
        if missing_source_notes:
            fail(f"notes lacking DOCX IDs: {missing_source_notes[:20]}", failures)

    # Section order.
    pos = []
    for title in REQUIRED_SECTION_TITLES:
        idx = next((i for i, t in enumerate(slide_texts, 1) if title in t), None)
        if idx is None:
            fail(f"missing section title: {title}", failures)
        else:
            pos.append(idx)
    if len(pos) == len(REQUIRED_SECTION_TITLES) and pos == sorted(pos):
        ok("required sections present in order")

    # Footer totals: slider3636 excludes title/section/key_message/transition from numbered slides.
    no_page_types = {"title", "section", "key_message", "transition"}
    expected_numbered = sum(1 for s in slides if s.get("type") not in no_page_types)
    totals = []
    badge_slides = []
    for i, text in enumerate(slide_texts, 1):
        badges = page_badges(text)
        if badges:
            badge_slides.append(i)
            totals.extend(int(t) for _, t in badges)
    if totals and all(t == expected_numbered for t in totals):
        ok(f"footer totals match numbered count {expected_numbered}")
    else:
        fail(f"footer totals mismatch: found {sorted(set(totals))}, expected {expected_numbered}", failures)

    # Density.
    dense = []
    too_dense = []
    for i, txt in enumerate(slide_texts, 1):
        if len(txt) > DENSE_CHARS:
            dense.append((i, len(txt)))
        if len(txt) > MAX_VISIBLE_CHARS:
            too_dense.append((i, len(txt)))
    if not too_dense and len(dense) <= MAX_DENSE_SLIDES:
        ok("visible text density acceptable")
    else:
        if too_dense:
            fail(f"too dense slides >{MAX_VISIBLE_CHARS}: {too_dense[:20]}", failures)
        if len(dense) > MAX_DENSE_SLIDES:
            fail(f"too many dense slides >{DENSE_CHARS}: {dense[:20]}", failures)

    # Sparse slide gate (Vietnamese, not structural).
    sparse_types = {"title", "section", "outline", "references"}
    thin = []
    for i, txt in enumerate(slide_texts, 1):
        stype = slides[i-1].get("type") if i <= len(slides) else ""
        if stype in sparse_types:
            continue
        core = re.sub(r"\b\d+\s+/\s+\d+\b", "", txt).strip()
        if len(core) < MIN_VISIBLE_CHARS:
            thin.append((i, stype, len(core)))
    if not thin:
        ok(f"no sparse content slides (>= {MIN_VISIBLE_CHARS} chars)")
    else:
        fail(f"sparse slides <{MIN_VISIBLE_CHARS}: {thin[:20]}", failures)

    # Final references/check there is no old footer denominator 149.
    if "149" not in "\n".join(re.findall(r"\d+\s*/\s*\d+", all_text)):
        ok("old /149 denominator absent")
    else:
        fail("old /149 denominator still present", failures)
    if "Tài liệu tham khảo" in slide_texts[-1]:
        ok("final slide is references")
    else:
        fail("final slide is not references", failures)

    print(f"text_chars: {len(all_text)}")
    print(f"picture_count: {pics}")
    print(f"footer_numbered_expected: {expected_numbered}")
    print(f"dense_slides: {dense[:20]}")

    if failures:
        print("\nRESULT: FAIL")
        raise SystemExit(1)
    print("\nRESULT: PASS")


if __name__ == "__main__":
    main()
