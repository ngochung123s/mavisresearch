# -*- coding: utf-8 -*-
"""Build the DOCX-only leiomyoma ART slider3636 PPTX and fail on count/provenance mismatch."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.stderr.reconfigure(encoding="utf-8", errors="replace")

from pptx import Presentation
from pptx.util import Emu, Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

ROOT = Path(__file__).parent
DATE = "2026-07-13"
ENGINE = Path("F:/DL/slider3636.py")
DECK = ROOT / f"U_xo_tu_cung_leiomyoma_slider3636_long_{DATE}.deck.json"
PROVENANCE = ROOT / f"U_xo_tu_cung_leiomyoma_slider3636_long_{DATE}.provenance.json"
OUT = ROOT / f"U_xo_tu_cung_leiomyoma_slider3636_long_smooth_{DATE}.pptx"
EXPECTED_SLIDES = 104
APPROVED_DOCX = ROOT / "U_xo_tu_cung_ART_Infertility_2026-07-12.docx"

FOOTER_LEFT = Inches(0.4)
FOOTER_TOP = Inches(7.08)
FOOTER_W = Inches(8.8)
FOOTER_H = Inches(0.3)

ALLOWED_WARNING_FIELDS = {"note", "cite"}


def load_slider():
    import types
    slider = types.ModuleType("slider3636")
    code_lines = []
    for line in ENGINE.read_text(encoding="utf-8").splitlines():
        if "Cell 5 · Theme Designer" in line:
            break
        stripped = line.lstrip()
        if stripped.startswith("!") or stripped.startswith("%"):
            continue
        code_lines.append(line)
    exec(compile("\n".join(code_lines), str(ENGINE), "exec"), slider.__dict__)
    return slider


def load_inputs():
    if not DECK.exists():
        raise SystemExit(f"Missing deck JSON: {DECK}")
    if not PROVENANCE.exists():
        raise SystemExit(f"Missing provenance JSON: {PROVENANCE}")
    deck = json.loads(DECK.read_text(encoding="utf-8"))
    provenance = json.loads(PROVENANCE.read_text(encoding="utf-8"))
    slides = deck.get("slides", [])
    prov_slides = provenance.get("slides", [])
    if len(slides) != EXPECTED_SLIDES:
        raise SystemExit(f"Deck has {len(slides)} slides; expected {EXPECTED_SLIDES}")
    if len(prov_slides) != len(slides):
        raise SystemExit(f"Provenance has {len(prov_slides)} entries; deck has {len(slides)}")
    src = Path(provenance.get("source_docx", ""))
    if src.resolve() != APPROVED_DOCX.resolve():
        raise SystemExit(f"Forbidden provenance source: {src}")
    for i, entry in enumerate(prov_slides, 1):
        if entry.get("slide_index") != i:
            raise SystemExit(f"Provenance index mismatch at {i}: {entry.get('slide_index')}")
        if Path(entry.get("source_docx", "")).resolve() != APPROVED_DOCX.resolve():
            raise SystemExit(f"Slide {i} has forbidden source: {entry.get('source_docx')}")
    return deck, provenance


def add_cite_footer(slide, pmids):
    if not pmids:
        return
    text = "Nguồn DOCX: " + " · ".join(f"PMID {p}" for p in sorted(set(pmids)))
    tb = slide.shapes.add_textbox(FOOTER_LEFT, FOOTER_TOP, FOOTER_W, FOOTER_H)
    tf = tb.text_frame
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    run.font.size = Pt(8)
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x72, 0x8C, 0x8C)


def patch_notes_and_citations(pptx_path: str | Path, deck: dict, provenance: dict):
    prs = Presentation(str(pptx_path))
    specs = deck["slides"]
    prov = provenance["slides"]
    if len(prs.slides) != len(specs):
        raise SystemExit(f"PPTX actual slide count {len(prs.slides)} != deck count {len(specs)}")
    for i, slide in enumerate(prs.slides):
        spec = specs[i]
        entry = prov[i]
        note_text = spec.get("note")
        if not note_text:
            if spec.get("type") in {"title", "outline", "section"}:
                note_text = f"Structural slide. Nguồn DOCX: {APPROVED_DOCX.name}."
            else:
                raise SystemExit(f"Slide {i+1} missing note")
        if spec.get("type") not in {"title", "outline", "section"} and "Nguồn DOCX:" not in note_text:
            raise SystemExit(f"Slide {i+1} note lacks DOCX provenance")
        slide.notes_slide.notes_text_frame.text = note_text
        if i > 0:
            add_cite_footer(slide, entry.get("pmids", []))
    prs.save(str(pptx_path))


def actual_slide_count(pptx_path: str | Path) -> int:
    return len(Presentation(str(pptx_path)).slides)


def validate_warnings(warnings: list[str]):
    bad = []
    for w in warnings:
        # slider3636 treats unknown fields as warning. For this workflow only note/cite warnings are acceptable.
        m = re.search(r"field không nhận diện: ([^()]+)", w)
        if not m:
            bad.append(w)
            continue
        fields = {x.strip() for x in m.group(1).split(",")}
        if not fields <= ALLOWED_WARNING_FIELDS:
            bad.append(w)
    if bad:
        print("UNACCEPTABLE WARNINGS:")
        for w in bad[:30]:
            print(" -", w)
        if len(bad) > 30:
            print(f" - ... {len(bad)-30} more")
        raise SystemExit(2)


def main():
    deck, provenance = load_inputs()
    slider = load_slider()
    slides = deck["slides"]
    errors, warnings = slider.validate_deck(slides)
    validate_warnings(warnings)
    if errors:
        print("ERRORS:")
        for e in errors:
            print(" -", e)
        raise SystemExit(2)

    import __main__
    class T: pass
    for k, v in slider.PALETTES["🟢 Medical Teal"].items():
        setattr(T, k, v)
    __main__.T = T()
    __main__.ACTIVE_FONT = "Arial"
    __main__.ACTIVE_FONT_SIZES = slider.FONT_SIZE_PRESETS["Large — hội trường lớn"].copy()
    __main__.SHOW_PAGE_NUM = True

    path, rendered = slider.build_presentation(deck, str(OUT))
    if rendered != EXPECTED_SLIDES:
        raise SystemExit(f"Rendered {rendered}/{EXPECTED_SLIDES}; refusing partial PPTX")
    if actual_slide_count(path) != EXPECTED_SLIDES:
        raise SystemExit(f"Actual PPTX slide count {actual_slide_count(path)} != {EXPECTED_SLIDES}")
    patch_notes_and_citations(path, deck, provenance)
    if actual_slide_count(path) != EXPECTED_SLIDES:
        raise SystemExit(f"Patched PPTX slide count changed: {actual_slide_count(path)}")
    print(f"OK: rendered {rendered}/{EXPECTED_SLIDES} slides -> {path}")
    print(f"Provenance: {PROVENANCE}")


if __name__ == "__main__":
    main()
