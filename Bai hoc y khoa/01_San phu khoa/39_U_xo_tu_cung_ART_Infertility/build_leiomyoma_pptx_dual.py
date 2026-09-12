# -*- coding: utf-8 -*-
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).parent
DATE = "2026-07-13"

SLIDES = [
    {
        "kind": "cover",
        "title": "U xơ cơ tử cung / Uterine leiomyoma",
        "subtitle": "Từ bệnh sinh, giải phẫu bệnh đến điều trị và ART",
        "tag": "Sản phụ khoa + Hỗ trợ sinh sản",
        "note": "Mục tiêu deck: biến bài học dài thành bản trình chiếu có logic làm gì và tại sao làm vậy.",
    },
    {
        "kind": "agenda",
        "title": "Bản đồ bài học",
        "items": ["1. Bản chất u xơ", "2. Dịch tễ và nguy cơ", "3. Bệnh sinh phân tử", "4. Giải phẫu bệnh", "5. FIGO và triệu chứng", "6. Chẩn đoán", "7. Xử trí chung", "8. Fertility/ART", "9. Theo dõi và tư vấn"],
        "note": "Làm gì? Đi từ bệnh học nền tới quyết định lâm sàng. Tại sao? Vì quyết định mổ phụ thuộc vị trí, triệu chứng và mục tiêu sinh sản, không chỉ kích thước.",
    },
    {
        "kind": "key",
        "title": "Câu hỏi đúng khi gặp u xơ",
        "lead": "Không phải: có u xơ thì mổ không?",
        "items": ["U nằm ở đâu?", "Có chạm/biến dạng khoang nội mạc không?", "Có gây AUB, thiếu máu, đau, chèn ép không?", "Can thiệp giúp hay làm chậm mục tiêu sinh sản?"],
        "note": "Làm gì? Đổi câu hỏi từ có/không mổ sang phân tầng. Tại sao? Vì u dưới thanh mạc nhỏ và u dưới niêm mạc nhỏ có ý nghĩa hoàn toàn khác nhau.",
    },
    {
        "kind": "two",
        "title": "U xơ là bệnh gì?",
        "left_title": "Định nghĩa",
        "left": ["Khối u lành tính của cơ trơn tử cung", "Thường đơn dòng, phụ thuộc hormone", "Có thể đơn độc hoặc nhiều khối"],
        "right_title": "Ý nghĩa thực hành",
        "right": ["Phát hiện tình cờ không đồng nghĩa cần điều trị", "Điều trị bệnh nhân, không điều trị hình ảnh", "Mục tiêu điều trị phải rõ"],
        "note": "Tại sao cần định nghĩa đúng? Vì gọi là u xơ dễ làm quên nguồn gốc cơ trơn và thành phần sinh học phức tạp của mô u.",
    },
    {
        "kind": "table",
        "title": "Yếu tố nguy cơ và modifier",
        "rows": [("Tuổi sinh sản", "Hoạt động dưới estrogen-progesterone"), ("Tiền sử gia đình", "Gợi ý nền di truyền"), ("Ít sinh/nulliparity", "Phơi nhiễm chu kỳ dài hơn"), ("Béo phì/tăng huyết áp", "Tăng nguy cơ và tăng rủi ro phẫu thuật"), ("Thai kỳ/mãn kinh", "Có thể thoái hóa đỏ; sau mãn kinh thường giảm")],
        "note": "Làm gì? Hỏi nguy cơ nhưng không điều trị dựa vào nguy cơ đơn thuần. Tại sao? Triệu chứng, khoang tử cung và mục tiêu bệnh nhân mới quyết định xử trí.",
    },
    {
        "kind": "mechanism",
        "title": "Bệnh sinh: không chỉ là một cục xơ",
        "steps": ["Tế bào cơ trơn đơn dòng", "Estrogen/progesterone", "TGF-beta, Wnt, hypoxia", "Tích lũy extracellular matrix", "Khối chắc, whorled, gây AUB/chèn ép"],
        "note": "Tại sao cần hiểu ECM? Vì thuốc có thể giảm triệu chứng/kích thước tạm thời nhưng không xóa hoàn toàn thành phần xơ và biến dạng cấu trúc.",
    },
    {
        "kind": "cards",
        "title": "Phân nhóm phân tử cần nhớ",
        "cards": [("MED12", "Phân nhóm lớn; không test thường quy"), ("HMGA2", "Một subtype sinh học khác"), ("FH-deficient/HLRCC", "Red flag nếu nhiều u, khởi phát sớm, mô học lạ")],
        "note": "Làm gì? Nhận diện ca bất thường cần hội chẩn GPB/di truyền. Tại sao? Hầu hết không cần xét nghiệm phân tử, nhưng FH-deficient là bối cảnh không nên bỏ sót.",
    },
    {
        "kind": "two",
        "title": "Đại thể và mô học điển hình",
        "left_title": "Đại thể",
        "left": ["Khối chắc, giới hạn rõ", "Trắng-ngà/trắng-xám", "Mặt cắt xoáy cuộn (whorled)", "Có thể nhiều khối"],
        "right_title": "Mô học",
        "right": ["Bó cơ trơn đan chéo", "Nhân hình thoi/cigar-shaped", "Ít dị dạng, ít phân bào", "Không có coagulative tumor cell necrosis"],
        "note": "Tại sao quan trọng? Vì mô học bland giúp phân biệt usual leiomyoma với STUMP/leiomyosarcoma trong bối cảnh có red flags.",
    },
    {
        "kind": "cards",
        "title": "Thoái hóa u xơ",
        "cards": [("Hyaline", "Thường gặp nhất"), ("Cystic/myxoid", "Có thể làm hình ảnh khó đọc"), ("Red degeneration", "Đau cấp, nhất là thai kỳ"), ("Calcific", "U lâu năm/sau mãn kinh")],
        "note": "Làm gì? Gắn kiểu thoái hóa với triệu chứng và hình ảnh. Tại sao? Đau cấp trong thai kỳ thường là red degeneration, không mặc định là xoắn hay ác tính.",
    },
    {
        "kind": "key",
        "title": "Red flags không được bỏ sót",
        "lead": "Đa số u xơ là lành tính, nhưng không gọi mọi khối cơ tử cung là lành nếu bối cảnh không phù hợp.",
        "items": ["Tăng nhanh sau mãn kinh", "Chảy máu sau mãn kinh", "MRI có hoại tử/xuất huyết/khuếch tán hạn chế đáng ngờ", "Bờ xâm lấn hoặc triệu chứng toàn thân", "Không morcellation không bảo vệ nếu nghi ngờ"],
        "note": "Tại sao? Imaging không loại trừ sarcoma 100%. Nếu nghi ngờ, chiến lược lấy bệnh phẩm phải thay đổi để tránh gieo rắc mô.",
    },
    {
        "kind": "table",
        "title": "FIGO type: vị trí quyết định ý nghĩa",
        "rows": [("0-2", "Dưới niêm mạc/lồi khoang — AUB, infertility, thường xử trí nếu mong thai"), ("3", "Trong cơ chạm nội mạc — vùng xám quan trọng trong ART"), ("4", "Trong cơ hoàn toàn — tùy kích thước/triệu chứng/khoang"), ("5-7", "Dưới thanh mạc — thường ít ảnh hưởng implantation nếu nhỏ"), ("8", "Cổ tử cung/dây chằng rộng/ký sinh — cá thể hóa")],
        "note": "Làm gì? Luôn yêu cầu FIGO type trong báo cáo. Tại sao? Một dòng 'u xơ tử cung' không đủ để quyết định IVF hay mổ.",
    },
    {
        "kind": "two",
        "title": "Triệu chứng: không chỉ vô sinh",
        "left_title": "Chảy máu",
        "left": ["AUB-L/rong kinh", "Thiếu máu thiếu sắt", "Ra huyết giữa kỳ nếu sát khoang"],
        "right_title": "Khối và đau",
        "right": ["Căng tức vùng chậu", "Tiểu lắt nhắt/bí tiểu", "Táo bón", "Đau cấp: red degeneration hoặc xoắn u có cuống"],
        "note": "Tại sao? Nếu chỉ nhìn qua ART sẽ bỏ sót chỉ định điều trị thật sự như thiếu máu hoặc chèn ép.",
    },
    {
        "kind": "algorithm",
        "title": "Chẩn đoán hình ảnh",
        "steps": ["TVUS 2D đầu tay", "3D TVUS/SIS nếu nghi tổn thương khoang", "Hysteroscopy khi cần chẩn đoán/điều trị trong khoang", "MRI khi đa u, u lớn, adenomyosis, sarcoma red flag hoặc mổ khó"],
        "note": "Làm gì? Chọn công cụ theo câu hỏi. Tại sao? HSG/TVUS mơ hồ không đủ khi quyết định chuyển phôi hoặc mổ type 0-3.",
    },
    {
        "kind": "table",
        "title": "Chẩn đoán phân biệt",
        "rows": [("Adenomyosis", "Đau bụng kinh, tử cung dày lan tỏa, giảm ART"), ("Polyp nội mạc", "Tổn thương mềm trong khoang"), ("Vách tử cung", "Sảy thai/vô sinh, cần 3D/MRI"), ("Khối buồng trứng", "Nhầm với u dưới thanh mạc có cuống"), ("Sarcoma/STUMP", "Hình ảnh/bối cảnh không điển hình")],
        "note": "Tại sao? Mổ u xơ mà bỏ sót adenomyosis hoặc polyp có thể không giải quyết nguyên nhân chính của triệu chứng/thất bại làm tổ.",
    },
    {
        "kind": "key",
        "title": "Nguyên tắc xử trí chung",
        "lead": "Điều trị mục tiêu, không điều trị hình ảnh.",
        "items": ["Không triệu chứng: theo dõi", "AUB/thiếu máu: kiểm soát chảy máu + sắt", "Đau/chèn ép: điều trị theo mức độ", "Mong thai/ART: ưu tiên khoang tử cung", "Không overpromise thuốc làm nhỏ u như cure cấu trúc"],
        "note": "Làm gì? Chốt mục tiêu trước khi chọn thuốc hay mổ. Tại sao? Cùng một u xơ có thể cần theo dõi, thuốc, myomectomy hoặc hysterectomy tùy bệnh nhân.",
    },
    {
        "kind": "table",
        "title": "Điều trị phụ khoa tổng quát",
        "rows": [("Theo dõi", "Không triệu chứng, không red flag"), ("NSAID/tranexamic acid", "Giảm đau/kinh nhiều, không làm nhỏ u"), ("COC/progestin/LNG-IUS", "Kiểm soát chảy máu ở ca chọn lọc"), ("GnRH agonist/antagonist", "Cầu nối, giảm chảy máu/kích thước"), ("Myomectomy/hysterectomy", "Bảo tồn tử cung hoặc điều trị dứt điểm")],
        "note": "Tại sao? Điều trị nội khoa chủ yếu kiểm soát triệu chứng; nếu u biến dạng khoang rõ thì vấn đề cấu trúc vẫn còn.",
    },
    {
        "kind": "cards",
        "title": "Can thiệp ít xâm lấn: tư vấn thận trọng",
        "cards": [("UAE", "Không thường quy khi đang tích cực mong thai"), ("RFA/TRFA", "Bằng chứng sinh sản đang phát triển"), ("HIFU/MRgFUS", "Có thể giảm thể tích; dữ liệu thai kỳ còn giới hạn")],
        "note": "Làm gì? Tách mục tiêu giảm triệu chứng khỏi tối ưu sinh sản. Tại sao? Một kỹ thuật tốt cho triệu chứng chưa chắc là lựa chọn tối ưu trước IVF/FET.",
    },
    {
        "kind": "algorithm",
        "title": "Fertility/ART algorithm",
        "steps": ["Mô tả số lượng, kích thước, FIGO, liên quan khoang", "Type 0-2 hoặc biến dạng khoang: xử trí trước chuyển phôi", "Type 3: cá thể hóa theo kích thước, diện tiếp xúc, thất bại chuyển phôi", "Type 4 lớn/nhiều/triệu chứng: cân nhắc mổ", "Type 5-7 nhỏ không triệu chứng: thường không mổ chỉ vì IVF"],
        "note": "Tại sao? Khoang tử cung là nơi chuyển phôi, nhưng phẫu thuật cũng có giá của nó: dính, sẹo, trì hoãn IVF.",
    },
    {
        "kind": "two",
        "title": "Type 0-2: cuộc chơi của buồng tử cung",
        "left_title": "Làm gì?",
        "left": ["Hysteroscopic myomectomy", "Cắt phần u lồi vào khoang", "Khôi phục bề mặt nội mạc", "Cân nhắc staged nếu type 2 sâu"],
        "right_title": "Tại sao?",
        "right": ["U chiếm chỗ nơi phôi bám", "AUB và viêm cục bộ", "Cắt quá sâu gây thủng/dính", "An toàn hơn lấy sạch bằng mọi giá"],
        "note": "Nếu bỏ qua type 0-2 trước FET có thể thất bại do nguyên nhân sửa được. Nếu cắt quá sâu có thể gây dính buồng tử cung.",
    },
    {
        "kind": "two",
        "title": "Type 3: vùng xám cần tôn trọng",
        "left_title": "Vì sao khó?",
        "left": ["Chạm nội mạc nhưng không lồi khoang", "Hysteroscopy có thể nhìn gần bình thường", "Có thể ảnh hưởng microenvironment"],
        "right_title": "Quyết định dựa vào",
        "right": ["Kích thước và diện tiếp xúc", "Số lần thất bại chuyển phôi", "Phôi euploid/chất lượng phôi", "Tuổi, AMH, nguy cơ phẫu thuật"],
        "note": "Làm gì? Không gộp type 3 với type 4 vô hại. Tại sao? Khoảng cách tới nội mạc có thể quyết định tác động lên receptivity.",
    },
    {
        "kind": "two",
        "title": "Type 4-7: cuộc chơi của ổ bụng",
        "left_title": "Làm gì?",
        "left": ["Laparoscopic/open/robotic myomectomy", "Rạch cơ trên u", "Bóc theo pseudocapsule", "Khâu tử cung nhiều lớp"],
        "right_title": "Tại sao?",
        "right": ["U nằm trong cơ/ngoài tử cung", "Không lấy tốt qua buồng tử cung", "Khâu quyết định sẹo tử cung", "Giảm chảy máu, tụ máu, vỡ tử cung sau này"],
        "note": "Lấy u là phá; khâu tử cung là xây lại. Nếu bệnh nhân còn muốn có thai, chất lượng khâu rất quan trọng.",
    },
    {
        "kind": "key",
        "title": "Hạn chế dính buồng tử cung",
        "lead": "Ít sang chấn hơn thường tốt hơn làm sạch triệt để.",
        "items": ["Bảo vệ lớp nội mạc đáy", "Tránh đốt rộng/sâu", "Không tạo hai diện thương đối diện nếu có thể", "Type 2 sâu: dừng đúng lúc, cân nhắc thì 2", "Barrier/balloon/estrogen/second-look ở ca nguy cơ cao"],
        "note": "Tại sao? Dính hình thành khi hai bề mặt thô áp sát nhau trong lúc lành thương. Nội mạc đáy là nguồn tái tạo nội mạc.",
    },
    {
        "kind": "key",
        "title": "Embryo banking trước myomectomy",
        "lead": "Ở bệnh nhân lớn tuổi/DOR, noãn là tài sản mất theo thời gian.",
        "items": ["Nếu u chưa cản chọc hút noãn", "Nếu chưa bắt buộc sửa khoang ngay", "Có thể tạo phôi trước", "Sau đó xử trí tử cung trước FET nếu cần"],
        "note": "Tại sao? Trì hoãn tạo phôi có thể hại hơn lợi ích lý thuyết của việc bóc u type 4 nhỏ không biến dạng khoang.",
    },
    {
        "kind": "table",
        "title": "Theo dõi sau can thiệp",
        "rows": [("Sau hysteroscopy", "Khoang phục hồi chưa? còn u/dính không?"), ("Sau mổ ổ bụng", "Sẹo cơ tử cung, thiếu máu, tái phát, kế hoạch thai kỳ"), ("Nếu mở khoang", "Ghi tường trình, tư vấn thai kỳ/sinh mổ"), ("Trước FET", "SIS/hysteroscopy nếu ca khó hoặc phôi quý")],
        "note": "Làm gì? Theo dõi theo rủi ro thủ thuật. Tại sao? Kết cục sinh sản phụ thuộc khoang lành và sẹo tử cung an toàn.",
    },
    {
        "kind": "cards",
        "title": "Clinical pearls",
        "cards": [("Không điều trị siêu âm", "Điều trị triệu chứng và mục tiêu"), ("FIGO bắt buộc", "Không có FIGO thì chưa đủ cho IVF"), ("Adenomyosis", "Đồng phạm hay bị bỏ sót"), ("UAE", "Không phải đường tắt cho bệnh nhân mong thai")],
        "note": "Tóm tắt các bẫy thực hành thường gặp khi đánh giá u xơ trong phụ khoa và ART.",
    },
    {
        "kind": "summary",
        "title": "Take-home messages",
        "items": ["U xơ quan trọng vì vị trí, triệu chứng và mục tiêu bệnh nhân.", "Type 0-2 và biến dạng khoang là nhóm tác động cao trước chuyển phôi.", "Type 3 cần được tách khỏi type 4.", "Type 5-7 nhỏ thường không mổ chỉ vì IVF.", "Myomectomy có lợi ích nhưng có giá: dính, sẹo, trì hoãn, tái phát.", "Trong buồng tử cung: bảo vệ nội mạc đáy quan trọng hơn lấy sạch bằng mọi giá."],
        "note": "Thông điệp cuối: mổ u xơ là cân bằng giữa lấy u và bảo vệ tử cung.",
    },
    {
        "kind": "references",
        "title": "Tài liệu tham khảo chính",
        "items": ["Mocanu et al. Fibroids and infertility. PMID 42104843; PMCID PMC13278667.", "FIGO/PALM-COEIN 2018. PMID 30198563.", "ACOG PB 228. PMID 34011888.", "ASRM guideline myoma/fertility. PMID 28865538.", "Intramural leiomyomas and fertility meta-analysis. PMID 38935974.", "FIGO type 3 IVF meta-analysis. PMID 37183601.", "Endometrial receptivity affected by fibroids. PMID 39625040."],
        "note": "Các citation chính dùng để định hướng, không đưa số liệu nhạy cảm nếu chưa đối chiếu full text.",
    },
]


def rgb(hexstr):
    hexstr = hexstr.strip('#')
    return RGBColor(int(hexstr[0:2], 16), int(hexstr[2:4], 16), int(hexstr[4:6], 16))


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def add_textbox(slide, text, x, y, w, h, size=18, color="1f2937", bold=False, align=None, font="Arial", margin=0.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    p = tf.paragraphs[0]
    p.text = text
    if align is not None:
        p.alignment = align
    for r in p.runs:
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = rgb(color)
    return box


def add_bullets(slide, items, x, y, w, h, size=17, color="1f2937", font="Arial"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.1)
    tf.margin_right = Inches(0.05)
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {item}"
        p.level = 0
        p.space_after = Pt(6)
        for r in p.runs:
            r.font.name = font
            r.font.size = Pt(size)
            r.font.color.rgb = rgb(color)
    return box


def add_round_rect(slide, x, y, w, h, fill, line=None, radius=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    shape.line.width = Pt(1)
    return shape


def add_notes(slide, text):
    if "tại sao" not in text.lower() and "why" not in text.lower():
        text += "\nTại sao? Giúp người học hiểu lý do lâm sàng phía sau thông điệp."
    slide.notes_slide.notes_text_frame.text = text


def page_badge(slide, idx, total, W, H, theme, font="Arial"):
    if idx == 1:
        return
    add_round_rect(slide, W - 0.92, H - 0.48, 0.62, 0.26, theme["badge_bg"], theme["badge_bg"])
    add_textbox(slide, f"{idx} / {total}", W - 0.88, H - 0.445, 0.54, 0.18, size=7.5, color=theme["badge_text"], bold=True, align=PP_ALIGN.CENTER, font=font, margin=0)


def header(slide, title, theme, W, font="Arial"):
    add_textbox(slide, title, 0.55, 0.28, W - 1.1, 0.42, size=21 if W <= 10.1 else 28, color=theme["title"], bold=True, font=font)
    add_round_rect(slide, 0.55, 0.78, 1.15, 0.035, theme["accent"], theme["accent"], radius=False)


def build_deck(path, style):
    prs = Presentation()
    if style == "canonical":
        W, H = 10.0, 5.625
        theme = {"bg":"F8FAFC", "title":"0F172A", "body":"1F2937", "muted":"64748B", "accent":"2563EB", "accent2":"0EA5E9", "card":"FFFFFF", "soft":"E0F2FE", "badge_bg":"DBEAFE", "badge_text":"1D4ED8", "dark":"0F172A"}
        font = "Arial"
    else:
        W, H = 13.33, 7.5
        theme = {"bg":"F6FBFA", "title":"0D2B2B", "body":"1A2F2F", "muted":"728C8C", "accent":"00977A", "accent2":"00B0F0", "card":"FFFFFF", "soft":"E8F8F5", "badge_bg":"00977A", "badge_text":"FFFFFF", "dark":"0D2B2B"}
        font = "Segoe UI"
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]
    total = len(SLIDES)

    for idx, spec in enumerate(SLIDES, 1):
        slide = prs.slides.add_slide(blank)
        set_bg(slide, theme["bg"])
        kind = spec["kind"]
        if kind == "cover":
            add_round_rect(slide, 0, 0, W, H, theme["dark"], theme["dark"], radius=False)
            add_round_rect(slide, 0.55, 0.55, 0.16, H - 1.1, theme["accent"], theme["accent"], radius=False)
            add_textbox(slide, spec["tag"], 0.9, 0.75, W - 1.8, 0.35, size=12 if W <= 10.1 else 15, color="A7F3D0" if style != "canonical" else "BFDBFE", bold=True, font=font)
            add_textbox(slide, spec["title"], 0.9, 1.55, W - 1.8, 1.25, size=32 if W <= 10.1 else 44, color="FFFFFF", bold=True, font=font)
            add_textbox(slide, spec["subtitle"], 0.92, 3.0, W - 2.0, 0.55, size=16 if W <= 10.1 else 22, color="E5E7EB", font=font)
            add_textbox(slide, f"Bác sĩ Ngọc Hưng • {DATE}", 0.92, H - 0.75, W - 2, 0.3, size=11 if W <= 10.1 else 14, color="CBD5E1", font=font)
        elif kind == "agenda":
            header(slide, spec["title"], theme, W, font)
            cols = 3
            col_w = (W - 1.3) / cols
            for i, item in enumerate(spec["items"]):
                x = 0.6 + (i % cols) * col_w
                y = 1.15 + (i // cols) * (0.72 if W <= 10.1 else 0.95)
                add_round_rect(slide, x, y, col_w - 0.18, 0.48 if W <= 10.1 else 0.65, theme["card"], "D8E4F0")
                add_textbox(slide, item, x + 0.12, y + 0.1, col_w - 0.42, 0.28 if W <= 10.1 else 0.42, size=11 if W <= 10.1 else 15, color=theme["body"], font=font)
        elif kind == "key":
            header(slide, spec["title"], theme, W, font)
            add_round_rect(slide, 0.7, 1.05, W - 1.4, 0.72 if W <= 10.1 else 0.9, theme["soft"], "BEE3F8")
            add_textbox(slide, spec["lead"], 0.9, 1.22, W - 1.8, 0.35 if W <= 10.1 else 0.5, size=15 if W <= 10.1 else 21, color=theme["title"], bold=True, font=font)
            add_bullets(slide, spec["items"], 0.95, 2.05, W - 1.9, H - 2.75, size=15 if W <= 10.1 else 20, color=theme["body"], font=font)
        elif kind == "two":
            header(slide, spec["title"], theme, W, font)
            cw = (W - 1.55) / 2
            for x, side, ttl in [(0.6, spec["left"], spec["left_title"]), (0.95 + cw, spec["right"], spec["right_title"])]:
                add_round_rect(slide, x, 1.08, cw, H - 1.65, theme["card"], "D8E4F0")
                add_textbox(slide, ttl, x + 0.22, 1.28, cw - 0.45, 0.32, size=14 if W <= 10.1 else 20, color=theme["accent"], bold=True, font=font)
                add_bullets(slide, side, x + 0.28, 1.78, cw - 0.55, H - 2.6, size=12.8 if W <= 10.1 else 17.5, color=theme["body"], font=font)
        elif kind == "table":
            header(slide, spec["title"], theme, W, font)
            y = 1.08
            row_h = 0.62 if W <= 10.1 else 0.82
            for a, b in spec["rows"]:
                add_round_rect(slide, 0.65, y, W - 1.3, row_h - 0.08, theme["card"], "D8E4F0")
                add_textbox(slide, a, 0.85, y + 0.12, 2.35 if W <= 10.1 else 3.1, 0.26, size=11.5 if W <= 10.1 else 16, color=theme["accent"], bold=True, font=font)
                add_textbox(slide, b, 3.05 if W <= 10.1 else 4.2, y + 0.12, W - (3.8 if W <= 10.1 else 5.0), 0.34 if W <= 10.1 else 0.48, size=11.2 if W <= 10.1 else 15.5, color=theme["body"], font=font)
                y += row_h
        elif kind == "mechanism" or kind == "algorithm":
            header(slide, spec["title"], theme, W, font)
            n = len(spec["steps"])
            box_w = (W - 1.3) / n
            y = 2.0 if W <= 10.1 else 2.75
            for i, step in enumerate(spec["steps"]):
                x = 0.65 + i * box_w
                add_round_rect(slide, x, y, box_w - 0.16, 0.85 if W <= 10.1 else 1.08, theme["card"], "D8E4F0")
                add_textbox(slide, str(i+1), x + 0.08, y - 0.34, 0.34, 0.28, size=13 if W <= 10.1 else 18, color=theme["accent"], bold=True, align=PP_ALIGN.CENTER, font=font)
                add_textbox(slide, step, x + 0.12, y + 0.16, box_w - 0.4, 0.55 if W <= 10.1 else 0.75, size=9.5 if W <= 10.1 else 13.5, color=theme["body"], bold=True, align=PP_ALIGN.CENTER, font=font)
                if i < n - 1:
                    add_textbox(slide, "→", x + box_w - 0.18, y + 0.28, 0.25, 0.2, size=18 if W <= 10.1 else 24, color=theme["accent"], bold=True, font=font)
        elif kind == "cards":
            header(slide, spec["title"], theme, W, font)
            count = len(spec["cards"])
            cols = 2 if count == 4 else 3
            card_w = (W - 1.4) / cols
            for i, (a, b) in enumerate(spec["cards"]):
                x = 0.65 + (i % cols) * card_w
                y = 1.28 + (i // cols) * (1.45 if W <= 10.1 else 1.9)
                add_round_rect(slide, x, y, card_w - 0.18, 1.02 if W <= 10.1 else 1.32, theme["card"], "D8E4F0")
                add_textbox(slide, a, x + 0.16, y + 0.15, card_w - 0.5, 0.28, size=13 if W <= 10.1 else 18, color=theme["accent"], bold=True, font=font)
                add_textbox(slide, b, x + 0.16, y + 0.52, card_w - 0.48, 0.36 if W <= 10.1 else 0.56, size=10.5 if W <= 10.1 else 14.5, color=theme["body"], font=font)
        elif kind == "summary" or kind == "references":
            header(slide, spec["title"], theme, W, font)
            add_bullets(slide, spec["items"], 0.85, 1.15, W - 1.7, H - 1.8, size=12.2 if W <= 10.1 else 16.5, color=theme["body"], font=font)
        add_notes(slide, spec.get("note", "Làm gì? Trình bày ý chính. Tại sao làm vậy? Giúp liên kết kiến thức với quyết định lâm sàng."))
        page_badge(slide, idx, total, W, H, theme, font)
    prs.save(path)


if __name__ == "__main__":
    build_deck(ROOT / f"U_xo_tu_cung_leiomyoma_canonical_{DATE}.pptx", "canonical")
    build_deck(ROOT / f"U_xo_tu_cung_leiomyoma_legacy_slider3636_style_{DATE}.pptx", "legacy")
    print("done")
