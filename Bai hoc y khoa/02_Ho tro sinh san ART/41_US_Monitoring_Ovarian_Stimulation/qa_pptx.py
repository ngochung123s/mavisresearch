import re
from pptx import Presentation
from pathlib import Path

p = Path(r"F:\DL\mavisresearch\Bai hoc y khoa\02_Ho tro sinh san ART\41_US_Monitoring_Ovarian_Stimulation\US_Monitoring_Ovarian_Stimulation_slider3636_v2_2026-07-04.pptx")
prs = Presentation(p)

allt = []
fonts = set()
for sl in prs.slides:
    for sh in sl.shapes:
        if hasattr(sh, "text") and sh.text:
            allt.append(sh.text)
        if hasattr(sh, "text_frame"):
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.name:
                        fonts.add(run.font.name)

full = "|".join(allt)
lower = full.lower()

pl = ["placeholder", "lorem", "ipsum"]
pmids = ["41732035", "32395637", "33280722", "21558332", "26597569"]
keys = ["AFC", "AMH", "antagonist", "GnRH", "freeze-all", "OHSS", "progesterone"]
bad_page = bool(re.search(r"(?m)^\s*\d{1,2}\s*/\s*\d{1,2}\s*$", full))

print(f"slides: {len(prs.slides)}")
print(f"chars: {len(full)}")
print(f"placeholder: {[w for w in pl if w in lower]}")
print(f"pmids: {{{', '.join(f'{p}:{p in full}' for p in pmids)}}}")
print(f"keys: {{{', '.join(f'{k}:{k.lower() in lower}' for k in keys)}}}")
print(f"fonts: {sorted(fonts)[:4]}")
print(f"page_number: {bad_page}")
