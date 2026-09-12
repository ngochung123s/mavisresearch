# -*- coding: utf-8 -*-
"""qa_from_md_pptx.py — QA gate cho PPTX render từ bài cô gửi.md.

Checks (mỗi dòng in OK/FAIL):
- PPTX tồn tại, mở được bằng python-pptx.
- Slide count nằm trong khoảng 25–60.
- Không placeholder/lorem/xxxx.
- Không page-number badge.
- Không footer signature (tác giả/chuyên khoa) trong slide visible.
- Không embedded pictures.
- Mỗi slide có speaker notes >= 30 chars và đủ 4 marker: Làm gì?, Tại sao?,
  Nếu bỏ qua/làm sai?, Nguồn MD:.
- Không rò label notes ra visible slide text.
- Density: mỗi slide visible <= 1200 chars; warn nếu 900–1200; cảnh báo nếu
  có slide non-reference > 1200.
- Key terms phải xuất hiện ít nhất 1 lần trong deck.
"""

import re
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).parent
PPTX = ROOT / "US_Monitoring_Ovarian_Stimulation_from_MD_2026-07-14.pptx"

MIN_SLIDES = 25
MAX_SLIDES = 60
MIN_NOTE_CHARS = 30
DENSE_CHARS = 900
MAX_VISIBLE_CHARS = 1200
MAX_DENSE_SLIDES = 12

KEY_TERMS = [
    "AFC", "AMH", "FSH", "LH", "antagonist", "PPOS",
    "hCG", "GnRH agonist", "trigger", "đông phôi toàn bộ",
    "niêm mạc tử cung", "14–24 mm", "34–36 giờ",
]

# Page-number badge: yêu cầu spaces quanh '/' để không false-positive
# kiểu "1/2 tổng số nang" hay "type 3/4".
PAGE_BADGE_RE = re.compile(r"\b\d+\s+\/\s+\d+\b")
FORBIDDEN_VISIBLE_RE = re.compile(
    r"(?i)\b(Làm gì:|Tại sao:|Nếu bỏ qua:|Nguồn MD:)\b"
)
PLACEHOLDER_RE = re.compile(r"(?i)\b(placeholder|lorem|ipsum|xxxx)\b")
SIGNATURE_TERMS = ["Bác sĩ Ngọc Hưng", "Bác sĩ", "ART / IVF-ICSI"]


def iter_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(sh.shapes)


def shape_text(sh):
    parts = []
    if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
        parts.append(sh.text_frame.text or "")
    if getattr(sh, "has_table", False) and sh.has_table:
        for row in sh.table.rows:
            for cell in row.cells:
                parts.append(cell.text or "")
    return "\n".join(p for p in parts if p)


def slide_visible_text(slide):
    return "\n".join(shape_text(sh) for sh in iter_shapes(slide.shapes))


def slide_notes(slide):
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
        return slide.notes_slide.notes_text_frame.text or ""
    return ""


def count_pictures(prs):
    n = 0
    for slide in prs.slides:
        for sh in iter_shapes(slide.shapes):
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                n += 1
    return n


def main():
    results = []

    def report(name, ok, detail=""):
        tag = "OK" if ok else "FAIL"
        line = f"[{tag}] {name}"
        if detail:
            line += f" — {detail}"
        results.append(ok)
        print(line)

    if not PPTX.exists():
        report("PPTX exists", False, f"missing: {PPTX}")
        sys.exit(1)

    prs = Presentation(PPTX)
    n_slides = len(prs.slides)
    report(
        "PPTX exists + opens",
        True,
        f"{n_slides} slides at {PPTX}",
    )

    report(
        "Slide count in range",
        MIN_SLIDES <= n_slides <= MAX_SLIDES,
        f"{n_slides} in [{MIN_SLIDES},{MAX_SLIDES}]",
    )

    visible_all = []
    notes_all = []
    for slide in prs.slides:
        visible_all.append(slide_visible_text(slide))
        notes_all.append(slide_notes(slide))

    full_visible = "\n".join(visible_all)
    full_notes = "\n".join(notes_all)
    full_lower = full_visible.lower()

    placeholders = PLACEHOLDER_RE.findall(full_visible)
    report(
        "No placeholder text",
        len(placeholders) == 0,
        f"hits={placeholders[:3]}",
    )

    page_badges = PAGE_BADGE_RE.findall(full_visible)
    report(
        "No page-number badges",
        len(page_badges) == 0,
        f"hits={page_badges[:3]}",
    )

    forbidden_vis = FORBIDDEN_VISIBLE_RE.findall(full_visible)
    report(
        "Notes labels not leaked to visible",
        len(forbidden_vis) == 0,
        f"hits={forbidden_vis[:3]}",
    )

    pic_count = count_pictures(prs)
    report(
        "No embedded pictures",
        pic_count == 0,
        f"picture_count={pic_count}",
    )

    sig_hits = [s for s in SIGNATURE_TERMS if s in full_visible]
    report(
        "No footer signature in visible",
        len(sig_hits) == 0,
        f"hits={sig_hits}",
    )

    # Notes completeness per slide
    bad_notes = []
    for i, note in enumerate(notes_all, start=1):
        if len(note.strip()) < MIN_NOTE_CHARS:
            bad_notes.append((i, "too short"))
            continue
        lowered = note.lower()
        missing = []
        for marker in ["làm gì?", "tại sao?", "nếu bỏ qua", "nguồn md"]:
            if marker not in lowered:
                missing.append(marker)
        if missing:
            bad_notes.append((i, f"missing={missing}"))
    report(
        "Speaker notes present and complete",
        len(bad_notes) == 0,
        f"bad_slides={bad_notes[:3]} (total bad={len(bad_notes)})",
    )

    # Density per slide (skip last slide if it is references)
    dense = []
    over = []
    for i, txt in enumerate(visible_all, start=1):
        core = PAGE_BADGE_RE.sub("", txt).strip()
        if i == n_slides and "Tài liệu" in txt:
            continue
        chars = len(core)
        if chars > MAX_VISIBLE_CHARS:
            over.append((i, chars))
        elif chars > DENSE_CHARS:
            dense.append((i, chars))

    report(
        "No slide > MAX_VISIBLE_CHARS",
        len(over) == 0,
        f"over={over[:3]} (total over={len(over)})",
    )
    report(
        "Limited dense slides 900-1200",
        len(dense) <= MAX_DENSE_SLIDES,
        f"dense={len(dense)} dense_examples={dense[:3]}",
    )

    # Key terms coverage
    missing_terms = [t for t in KEY_TERMS if t.lower() not in full_lower]
    report(
        "Key terms present",
        len(missing_terms) == 0,
        f"missing={missing_terms}",
    )

    print()
    if all(results):
        print(f"RESULT: PASS — {sum(results)}/{len(results)} checks OK")
        sys.exit(0)
    print(f"RESULT: FAIL — {sum(results)}/{len(results)} checks OK")
    sys.exit(1)


if __name__ == "__main__":
    main()
