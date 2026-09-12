import re
from pathlib import Path

from pptx import Presentation


PPTX = Path(__file__).parent / "Sieu_am_theo_doi_kich_thich_buong_trung_slider3636_2026-07-03.pptx"
PMIDS = ["32395637", "41732035", "33280722", "21558332", "26597569"]
KEY_TERMS = ["AFC", "AMH", "antagonist", "GnRH", "freeze-all", "OHSS", "progesterone", "15", "noãn"]
PLACEHOLDERS = ["placeholder", "lorem", "ipsum", "xxxx"]


def text_of(slide):
    parts = []
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text:
            parts.append(shape.text)
    return "\n".join(parts)


def main():
    prs = Presentation(PPTX)
    text = "\n".join(text_of(slide) for slide in prs.slides)
    fonts = set()
    for slide in prs.slides:
        for shape in slide.shapes:
            if not hasattr(shape, "text_frame"):
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.name:
                        fonts.add(run.font.name)
    lower = text.lower()
    placeholder_hits = [w for w in PLACEHOLDERS if w in lower]
    pmids = {p: p in text for p in PMIDS}
    keys = {k: k.lower() in lower for k in KEY_TERMS}
    page_like = bool(re.search(r"(?m)^\s*\d{1,2}\s*/\s*\d{1,2}\s*$", text))
    print(f"PPTX: {PPTX}")
    print(f"slides: {len(prs.slides)}")
    print(f"text_chars: {len(text)}")
    print(f"placeholder_hits: {placeholder_hits}")
    print(f"pmids: {pmids}")
    print(f"key_terms: {keys}")
    print(f"fonts_sample: {sorted(fonts)[:8]}")
    print(f"page_number_like: {page_like}")
    failed = bool(placeholder_hits) or not all(pmids.values()) or not all(keys.values()) or page_like
    raise SystemExit(1 if failed else 0)


if __name__ == "__main__":
    main()
