from pathlib import Path

from pptx import Presentation


PPTX = Path(__file__).parent / "slides" / "output" / "Sieu_am_theo_doi_kich_thich_buong_trung_PPTXGenJS_2026-07-03.pptx"
PMIDS = ["32395637", "41732035", "33280722", "21558332", "26597569"]
KEY_TERMS = [
    "ovarian reserve",
    "ovarian response",
    "AFC",
    "AMH",
    "antagonist",
    "GnRH agonist trigger",
    "freeze-all",
    "OHSS",
    "progesterone",
    "15 noãn",
]
PLACEHOLDERS = ["placeholder", "lorem", "ipsum", "xxxx"]


def slide_text(slide):
    parts = []
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text:
            parts.append(shape.text)
    return "\n".join(parts)


def main():
    prs = Presentation(PPTX)
    texts = [slide_text(slide) for slide in prs.slides]
    all_text = "\n".join(texts)
    lower = all_text.lower()

    placeholder_hits = [word for word in PLACEHOLDERS if word in lower]
    pmid_status = {pmid: pmid in all_text for pmid in PMIDS}
    key_status = {term: term.lower() in lower for term in KEY_TERMS}

    print(f"PPTX: {PPTX}")
    print(f"slides: {len(prs.slides)}")
    print(f"text_chars: {len(all_text)}")
    print(f"placeholder_hits: {placeholder_hits}")
    print(f"pmids: {pmid_status}")
    print(f"key_terms: {key_status}")
    print("page_badge_check: skipped by user request")
    print(f"last_slide_has_refs: {'Tài liệu tham khảo' in texts[-1] and 'PMID' in texts[-1]}")

    failed = bool(placeholder_hits) or not all(pmid_status.values()) or not all(key_status.values())
    raise SystemExit(1 if failed else 0)


if __name__ == "__main__":
    main()
