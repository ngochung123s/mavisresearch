# -*- coding: utf-8 -*-
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.stderr.reconfigure(encoding="utf-8", errors="replace")

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).parent
DATE = "2026-07-13"
PPTX = ROOT / f"Semen_Analysis_Male_Infertility_slider3636_{DATE}.pptx"
PLACEHOLDERS = ["placeholder", "lorem", "ipsum", "xxxx"]
KEY_TERMS = [
    "WHO 2021", "AUA/ASRM", "EAU", "WHO manual", "2-7 ngày", "30 phút", "60 phút",
    "1.4 mL", "16 triệu/mL", "39 triệu/ejaculate", "42%", "30%", "54%", "4%",
    "lower fifth-centile", "không phải cutoff", "fertile/infertile", "pellet", "repeat",
    "3.000g", "vitality", "morphology", "SDF", "ROS", "ASA", "culture", "TMSC",
    "CBAVD", "EOD", "NOA", "OA", "Y-chromosome", "micro-TESE", "ICSI", "IUI",
    "cryptozoospermia", "varicocele", "round cells", "leukocytospermia", "agglutination", "aggregation", "bài 42",
]
PMIDS_OR_IDS = ["39145501", "33528873", "35849333", "31400948", "9789240030787"]
MIN_SLIDES = 55
MAX_NON_REF_CHARS = 1200
DENSE_CHARS = 950
MAX_DENSE_SLIDES = 8


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def shape_text(shape):
    parts = []
    if getattr(shape, "has_text_frame", False):
        parts.append(shape.text_frame.text or "")
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                parts.append(cell.text or "")
    return "\n".join(p for p in parts if p)


def slide_text(slide):
    return "\n".join(shape_text(s) for s in iter_shapes(slide.shapes) if shape_text(s))


def notes_text(slide):
    if not slide.has_notes_slide:
        return ""
    return (slide.notes_slide.notes_text_frame.text or "").strip()


def main():
    prs = Presentation(PPTX)
    slide_texts = [slide_text(slide).strip() for slide in prs.slides]
    all_text = "\n".join(slide_texts)
    lower = all_text.lower()
    note_texts = [notes_text(slide) for slide in prs.slides]
    thin_notes = [i + 1 for i, n in enumerate(note_texts) if len(n) < 40 or "tại sao" not in n.lower()]
    placeholder_hits = [w for w in PLACEHOLDERS if w in lower]
    missing_terms = [t for t in KEY_TERMS if t.lower() not in lower]
    missing_ids = [p for p in PMIDS_OR_IDS if p not in all_text]
    forbidden_labels = ["Làm gì:", "Tại sao:", "Bằng chứng:", "Nếu bỏ qua/làm sai:", "Nguy cơ nếu sai:"]
    forbidden_hits = [label for label in forbidden_labels if label in all_text]
    dense_slides = []
    too_dense_slides = []
    for i, text in enumerate(slide_texts, 1):
        is_ref = "Tài liệu" in text or text.count("PMID") >= 3
        if not is_ref and len(text) > DENSE_CHARS:
            dense_slides.append((i, len(text)))
        if not is_ref and len(text) > MAX_NON_REF_CHARS:
            too_dense_slides.append((i, len(text)))
    substance_counts = {
        "WHO": all_text.count("WHO"),
        "không": lower.count("không"),
        "repeat": lower.count("repeat"),
        "pellet": lower.count("pellet"),
        "ART": all_text.count("ART"),
        "ICSI": all_text.count("ICSI"),
        "TMSC": all_text.count("TMSC"),
    }
    slide_lengths = sorted(len(t) for t in slide_texts)
    median_visible_chars = slide_lengths[len(slide_lengths) // 2]
    weak_substance = substance_counts["WHO"] < 10 or substance_counts["pellet"] < 4 or substance_counts["ICSI"] < 6
    sparse_visible = median_visible_chars < 160
    too_many_dense = len(dense_slides) > MAX_DENSE_SLIDES
    print(f"PPTX: {PPTX}")
    print(f"slides: {len(prs.slides)}")
    print(f"text_chars: {len(all_text)}")
    print(f"median_visible_chars: {median_visible_chars}")
    print(f"substance_counts: {substance_counts}")
    print(f"forbidden_label_hits: {forbidden_hits}")
    print(f"placeholder_hits: {placeholder_hits}")
    print(f"missing_terms: {missing_terms}")
    print(f"missing_pmids_or_ids: {missing_ids}")
    print(f"dense_slides_over_{DENSE_CHARS}: {dense_slides}")
    print(f"too_dense_slides_over_{MAX_NON_REF_CHARS}: {too_dense_slides}")
    print(f"thin_or_no_why_notes: {thin_notes[:20]}{'...' if len(thin_notes) > 20 else ''}")
    failed = (
        len(prs.slides) < MIN_SLIDES
        or bool(placeholder_hits)
        or bool(missing_terms)
        or bool(missing_ids)
        or bool(thin_notes)
        or bool(forbidden_hits)
        or bool(too_dense_slides)
        or too_many_dense
        or sparse_visible
        or weak_substance
    )
    raise SystemExit(1 if failed else 0)


if __name__ == "__main__":
    main()
