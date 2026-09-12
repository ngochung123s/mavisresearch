# -*- coding: utf-8 -*-
import importlib.util
import json
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.stderr.reconfigure(encoding="utf-8", errors="replace")

from pptx import Presentation

ROOT = Path(__file__).parent
DATE = "2026-07-13"
STEM = "Semen_Analysis_Male_Infertility"
ENGINE = Path("F:/DL/mavisresearch/Bai hoc y khoa/10_Script Python/slider3636.py")
DECK = ROOT / f"{STEM}_slider3636_{DATE}.deck.json"
OUT = ROOT / f"{STEM}_slider3636_{DATE}.pptx"
DEFAULT_NOTE = "Làm gì? Nhấn mạnh ý chính của phần này.\nTại sao? Giúp người học đặt nội dung vào quyết định lâm sàng.\nNếu bỏ qua/làm sai? Dễ áp dụng máy móc một con số."


def load_slider():
    import types
    slider = types.ModuleType("slider3636")
    code_lines = []
    for line in ENGINE.read_text(encoding="utf-8").splitlines():
        if "Cell 5 · Theme Designer" in line:
            break
        stripped = line.lstrip()
        if stripped.startswith("!") or stripped.startswith("%"):
            continue
        code_lines.append(line)
    exec(compile("\n".join(code_lines), str(ENGINE), "exec"), slider.__dict__)
    return slider


def patch_notes(pptx_path, deck):
    prs = Presentation(pptx_path)
    specs = deck.get("slides", [])
    for i, slide in enumerate(prs.slides):
        spec_note = specs[i].get("note") if i < len(specs) else ""
        text = spec_note or DEFAULT_NOTE
        if "tại sao" not in text.lower():
            text += "\nTại sao? Giúp liên hệ kiến thức với xử trí bệnh nhân."
        slide.notes_slide.notes_text_frame.text = text
    prs.save(pptx_path)


def main():
    slider = load_slider()
    deck = json.loads(DECK.read_text(encoding="utf-8"))
    slides = deck.get("slides", deck if isinstance(deck, list) else [])
    errors, warnings = slider.validate_deck(slides)
    if warnings:
        print("WARNINGS:")
        for w in warnings[:20]:
            print(" -", w)
        if len(warnings) > 20:
            print(f" - ... {len(warnings) - 20} more")
    if errors:
        print("ERRORS:")
        for e in errors:
            print(" -", e)
        raise SystemExit(2)
    import __main__
    class T: pass
    for k, v in slider.PALETTES["🟢 Medical Teal"].items():
        setattr(T, k, v)
    __main__.T = T()
    __main__.ACTIVE_FONT = "Arial"
    __main__.ACTIVE_FONT_SIZES = slider.FONT_SIZE_PRESETS["Large — hội trường lớn"].copy()
    __main__.SHOW_PAGE_NUM = True
    path, count = slider.build_presentation(deck, str(OUT))
    patch_notes(path, deck)
    print(f"OK: {count} slides -> {path}")


if __name__ == "__main__":
    main()
